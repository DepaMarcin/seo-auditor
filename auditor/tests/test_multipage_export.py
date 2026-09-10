"""Testy audytu wielu szablonów podstron oraz eksportu raportu.

Skanowanie podstron, parsowanie sitemap.xml i Google Sheets API są mockowane -
żaden test nie łączy się z siecią.
"""
from __future__ import annotations

import csv
import io
from unittest.mock import MagicMock, patch

import httpx
from django.contrib.auth.models import User
from django.core.cache import cache
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from auditor.models import Audit, AuditedPage, AuditMetric
from auditor.services.exporter import build_report, report_filename
from auditor.services.sheets import (
    GoogleSheetsService,
    MissingSheetsScopeError,
    SheetsExportError,
)
from auditor.services.sitemap import SitemapService
from auditor.services.spreadsheet import build_csv, build_xlsx


def _mock_client(get_return=None, get_side_effect=None):
    client = MagicMock()
    client.__enter__.return_value = client
    client.__exit__.return_value = False
    if get_side_effect is not None:
        client.get.side_effect = get_side_effect
    else:
        client.get.return_value = get_return
    return client


class AuditedPageModelTests(TestCase):
    def test_status_counts_summarises_metrics(self):
        audit = Audit.objects.create(url="https://example.com")
        page = AuditedPage.objects.create(
            audit=audit, url="https://example.com/produkt/x",
            page_type=AuditedPage.PageType.PRODUCT,
            metrics_data=[
                {"status": "error"}, {"status": "error"},
                {"status": "warning"}, {"status": "ok"}, {"status": "info"},
            ],
        )

        self.assertEqual(page.status_counts, {"error": 2, "warning": 1, "ok": 1, "info": 1})

    def test_same_url_cannot_repeat_within_one_audit(self):
        from django.db import IntegrityError

        audit = Audit.objects.create(url="https://example.com")
        AuditedPage.objects.create(audit=audit, url="https://example.com/a")

        with self.assertRaises(IntegrityError):
            AuditedPage.objects.create(audit=audit, url="https://example.com/a")


@patch("auditor.views.validate_public_url", side_effect=lambda url: url or "")
@patch("auditor.views.enqueue_audit")
class TemplateFormTests(TestCase):
    """Formularz startowy zapisuje dodatkowe szablony jako `AuditedPage`."""

    def setUp(self):
        cache.clear()
        self.user = User.objects.create_user(username="szablony", password="haslo12345")
        self.client.force_login(self.user)

    def test_additional_templates_are_saved_with_page_types(self, _enqueue, _guard):
        self.client.post(reverse("auditor:index"), {
            "url": "https://example.com",
            "url_category": "https://example.com/kategoria/buty",
            "url_product": "https://example.com/produkt/but",
            "url_blog": "https://example.com/blog/wpis",
            "url_offer": "https://example.com/oferta",
        })

        audit = Audit.objects.get()
        pages = {p.page_type: p.url for p in audit.pages.all()}

        self.assertEqual(len(pages), 4)
        self.assertEqual(pages[AuditedPage.PageType.CATEGORY], "https://example.com/kategoria/buty")
        self.assertEqual(pages[AuditedPage.PageType.BLOG], "https://example.com/blog/wpis")

    def test_empty_template_fields_are_skipped(self, _enqueue, _guard):
        self.client.post(reverse("auditor:index"), {
            "url": "https://example.com",
            "url_category": "https://example.com/kategoria",
            "url_product": "",
            "url_blog": "   ",
        })

        self.assertEqual(Audit.objects.get().pages.count(), 1)

    def test_duplicate_of_primary_url_is_ignored(self, _enqueue, _guard):
        """Adres główny i tak trafi do zestawienia - powtórzenie łamałoby unikalność."""
        self.client.post(reverse("auditor:index"), {
            "url": "https://example.com",
            "url_category": "https://example.com/",
        })

        self.assertEqual(Audit.objects.get().pages.count(), 0)

    def test_unsafe_template_url_blocks_the_whole_audit(self, _enqueue, guard):
        from auditor.services.url_guard import UnsafeUrlError

        guard.side_effect = lambda url: (_ for _ in ()).throw(
            UnsafeUrlError("Adresy w sieci lokalnej nie podlegają audytowi.")
        ) if "127.0.0.1" in url else url

        response = self.client.post(reverse("auditor:index"), {
            "url": "https://example.com",
            "url_category": "http://127.0.0.1:8000/admin/",
        })

        self.assertEqual(Audit.objects.count(), 0)
        self.assertRedirects(response, reverse("auditor:index"))


class ScanAdditionalPagesTests(TestCase):
    """Skanowanie szablonów: wynik per podstrona i odporność na błąd pojedynczej strony."""

    def setUp(self):
        from auditor.tests.test_audit_service import _service_with_mocks

        self.audit = Audit.objects.create(url="https://example.com")
        self.service = _service_with_mocks()

    def test_each_page_gets_its_own_metrics_and_score(self):
        AuditedPage.objects.create(
            audit=self.audit, url="https://example.com/produkt/x",
            page_type=AuditedPage.PageType.PRODUCT,
        )

        self.service.run_audit(self.audit)

        page = self.audit.pages.get(url="https://example.com/produkt/x")
        self.assertEqual(page.status, AuditedPage.Status.COMPLETED)
        self.assertGreater(len(page.metrics_data), 0)
        self.assertGreater(page.score, 0)

    def test_primary_url_is_stored_as_homepage_template(self):
        self.service.run_audit(self.audit)

        homepage = self.audit.pages.get(url=self.audit.url)
        self.assertEqual(homepage.page_type, AuditedPage.PageType.HOMEPAGE)
        self.assertEqual(homepage.status, AuditedPage.Status.COMPLETED)

    def test_failed_page_does_not_break_the_audit(self):
        from auditor.services.scraper import ScraperError

        AuditedPage.objects.create(audit=self.audit, url="https://example.com/zepsuta")

        def scrape(url):
            if "zepsuta" in url:
                raise ScraperError("404")
            return dict(self.service.scraper.scrape.return_value)

        self.service.scraper.scrape.side_effect = scrape

        self.service.run_audit(self.audit)

        self.audit.refresh_from_db()
        self.assertEqual(self.audit.status, Audit.Status.COMPLETED)
        broken = self.audit.pages.get(url="https://example.com/zepsuta")
        self.assertEqual(broken.status, AuditedPage.Status.FAILED)
        self.assertIn("Nie udało się pobrać", broken.error_message)

    def test_page_scan_skips_ai_recommendations(self):
        """Rekomendacje AI powstają raz, dla adresu głównego - nie per szablon."""
        AuditedPage.objects.create(audit=self.audit, url="https://example.com/blog/wpis")

        self.service.run_audit(self.audit)

        page = self.audit.pages.get(url="https://example.com/blog/wpis")
        for metric in page.metrics_data:
            self.assertNotIn("recommendation", metric["value"])


class SitemapServiceTests(SimpleTestCase):
    SITEMAP = """<?xml version="1.0" encoding="UTF-8"?>
    <urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">
        <url><loc>https://example.com/</loc></url>
        <url><loc>https://example.com/kategoria/buty</loc></url>
        <url><loc>https://example.com/produkt/but-sportowy</loc></url>
        <url><loc>https://example.com/blog/jak-dobrac-buty</loc></url>
        <url><loc>https://example.com/oferta/wspolpraca</loc></url>
    </urlset>"""

    def setUp(self):
        cache.clear()
        self.service = SitemapService()

    def _response(self, body: str, status: int = 200) -> httpx.Response:
        return httpx.Response(status, request=httpx.Request("GET", "https://example.com/sitemap.xml"), text=body)

    def test_classifies_one_url_per_template_type(self):
        with patch("auditor.services.sitemap.httpx.get", return_value=self._response(self.SITEMAP)), \
                patch("auditor.services.sitemap.validate_public_url", side_effect=lambda u: u):
            result = self.service.suggest_pages("https://example.com")

        self.assertTrue(result["available"])
        suggestions = result["suggestions"]
        self.assertEqual(suggestions["category"], "https://example.com/kategoria/buty")
        self.assertEqual(suggestions["product"], "https://example.com/produkt/but-sportowy")
        self.assertEqual(suggestions["blog"], "https://example.com/blog/jak-dobrac-buty")
        self.assertEqual(suggestions["offer"], "https://example.com/oferta/wspolpraca")

    def test_sitemap_index_is_followed_one_level_down(self):
        index = """<?xml version="1.0" encoding="UTF-8"?>
        <sitemapindex xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">
            <sitemap><loc>https://example.com/sitemap-1.xml</loc></sitemap>
        </sitemapindex>"""

        responses = [self._response("", 404), self._response(index), self._response(self.SITEMAP)]
        with patch("auditor.services.sitemap.httpx.get", side_effect=responses), \
                patch("auditor.services.sitemap.validate_public_url", side_effect=lambda u: u):
            result = self.service.suggest_pages("https://example.com")

        self.assertTrue(result["available"])
        self.assertIn("product", result["suggestions"])

    def test_missing_sitemap_is_not_an_error(self):
        with patch("auditor.services.sitemap.httpx.get", return_value=self._response("", 404)), \
                patch("auditor.services.sitemap.validate_public_url", side_effect=lambda u: u):
            result = self.service.suggest_pages("https://example.com")

        self.assertFalse(result["available"])
        self.assertEqual(result["suggestions"], {})
        self.assertIn("Nie znaleziono mapy witryny", result["error"])

    def test_malformed_xml_does_not_raise(self):
        with patch("auditor.services.sitemap.httpx.get", return_value=self._response("<to nie jest xml")), \
                patch("auditor.services.sitemap.validate_public_url", side_effect=lambda u: u):
            result = self.service.suggest_pages("https://example.com")

        self.assertFalse(result["available"])

    def test_unsafe_domain_is_rejected(self):
        result = self.service.suggest_pages("http://127.0.0.1:8000/")

        self.assertFalse(result["available"])


class ReportBuilderTests(TestCase):
    """Struktura raportu: trzy zakładki o ustalonych kolumnach."""

    def setUp(self):
        self.audit = Audit.objects.create(
            url="https://example.com", status=Audit.Status.COMPLETED, score=64,
            gsc_total_clicks_current=1500, gsc_total_clicks_previous=1200,
            gsc_yoy_change_percent=25.0,
            gsc_top_gainers=[{
                "query": "audyt seo", "clicks_current": 100, "clicks_previous": 40, "delta": 60,
                "impressions_current": 5000, "ctr_current": 2.0, "position_current": 8.4,
            }],
            ga4_organic_sessions=401, senuto_top10=42,
        )
        AuditMetric.objects.create(
            audit=self.audit, category="seo", key="meta_description", status="error",
            value={"note": "Brak meta description.", "recommendation": "Dodaj meta description."},
        )
        AuditMetric.objects.create(
            audit=self.audit, category="technical", key="canonical", status="warning",
            value={"note": "Brak canonical."},
        )
        AuditedPage.objects.create(
            audit=self.audit, url="https://example.com", page_type=AuditedPage.PageType.HOMEPAGE,
            status=AuditedPage.Status.COMPLETED,
            metrics_data=[{"category": "seo", "key": "meta_description", "status": "error",
                           "value": {"note": "Brak meta description."}, "current_value": ""}],
        )
        AuditedPage.objects.create(
            audit=self.audit, url="https://example.com/produkt/x", page_type=AuditedPage.PageType.PRODUCT,
            status=AuditedPage.Status.COMPLETED,
            metrics_data=[{"category": "seo", "key": "title", "status": "ok",
                           "value": {"note": "Tytuł poprawny."}, "current_value": ""}],
        )

    def test_report_has_three_named_sheets(self):
        sheets = build_report(self.audit)

        self.assertEqual(
            [s.title for s in sheets],
            ["Ruch i Widoczność", "Matryca Techniczna Szablonów", "Priorytetowy Backlog Zadań"],
        )

    def test_traffic_sheet_includes_gsc_metrics(self):
        traffic = build_report(self.audit)[0]
        joined = [str(cell) for row in traffic.rows for cell in row]

        self.assertIn("Wyświetlenia", traffic.headers)
        self.assertIn("CTR %", traffic.headers)
        self.assertIn("Śr. pozycja", traffic.headers)
        self.assertIn("audyt seo", joined)
        self.assertIn("5000", joined)

    def test_matrix_sheet_has_required_columns_and_covers_all_templates(self):
        matrix = build_report(self.audit)[1]

        self.assertEqual(
            matrix.headers,
            ["Typ Szablonu", "Adres URL", "Kategoria", "Metryka", "Status", "Diagnoza", "Rekomendacja AI"],
        )
        page_types = {row[0] for row in matrix.rows}
        self.assertIn("Strona główna", page_types)
        self.assertIn("Strona produktu", page_types)

    def test_matrix_reuses_main_audit_recommendation_for_templates(self):
        """Podstrony nie mają własnych rekomendacji - biorą je z audytu głównego."""
        matrix = build_report(self.audit)[1]
        homepage_row = next(r for r in matrix.rows if r[0] == "Strona główna")

        self.assertEqual(homepage_row[6], "Dodaj meta description.")

    def test_failed_page_is_reported_in_matrix(self):
        AuditedPage.objects.create(
            audit=self.audit, url="https://example.com/zepsuta",
            status=AuditedPage.Status.FAILED, error_message="Nie udało się pobrać podstrony: 404",
        )

        matrix = build_report(self.audit)[1]
        row = next(r for r in matrix.rows if r[1] == "https://example.com/zepsuta")

        self.assertEqual(row[4], "BŁĄD")
        self.assertIn("404", row[5])

    def test_backlog_lists_only_problems_sorted_by_priority(self):
        backlog = build_report(self.audit)[2]

        statuses = {row[4] for row in backlog.rows}
        self.assertTrue(statuses <= {"BŁĄD", "OSTRZEŻENIE"})
        priorities = [row[0] for row in backlog.rows]
        self.assertEqual(priorities, sorted(priorities, reverse=True))

    def test_backlog_shows_placeholder_when_no_problems(self):
        self.audit.metrics.all().delete()

        backlog = build_report(self.audit)[2]

        self.assertEqual(len(backlog.rows), 1)
        self.assertIn("Brak wykrytych błędów", backlog.rows[0][3])

    def test_audit_without_pages_still_produces_matrix(self):
        """Audyty sprzed wprowadzenia szablonów nie mają rekordów AuditedPage."""
        self.audit.pages.all().delete()

        matrix = build_report(self.audit)[1]

        self.assertGreater(len(matrix.rows), 0)
        self.assertEqual({row[0] for row in matrix.rows}, {"Strona główna"})

    def test_filename_contains_domain_and_date(self):
        name = report_filename(self.audit, "xlsx")

        self.assertIn("example.com", name)
        self.assertTrue(name.endswith(".xlsx"))


class SpreadsheetFormatTests(TestCase):
    def setUp(self):
        self.audit = Audit.objects.create(url="https://example.com", status=Audit.Status.COMPLETED)
        AuditMetric.objects.create(
            audit=self.audit, category="seo", key="title", status="error",
            value={"note": "Brak tytułu."},
        )

    def test_xlsx_has_one_worksheet_per_report_sheet(self):
        workbook = load_workbook(io.BytesIO(build_xlsx(build_report(self.audit))))

        self.assertEqual(
            workbook.sheetnames,
            ["Ruch i Widoczność", "Matryca Techniczna Szablonów", "Priorytetowy Backlog Zadań"],
        )
        self.assertEqual(workbook["Priorytetowy Backlog Zadań"].freeze_panes, "A2")

    def test_xlsx_headers_match_report_structure(self):
        workbook = load_workbook(io.BytesIO(build_xlsx(build_report(self.audit))))
        worksheet = workbook["Matryca Techniczna Szablonów"]

        headers = [cell.value for cell in worksheet[1]]
        self.assertEqual(headers[0], "Typ Szablonu")
        self.assertEqual(headers[-1], "Rekomendacja AI")

    def test_csv_contains_all_three_sections(self):
        content = build_csv(build_report(self.audit))

        self.assertIn("### Ruch i Widoczność ###", content)
        self.assertIn("### Matryca Techniczna Szablonów ###", content)
        self.assertIn("### Priorytetowy Backlog Zadań ###", content)

    def test_csv_is_parsable_with_semicolon_delimiter(self):
        rows = list(csv.reader(io.StringIO(build_csv(build_report(self.audit))), delimiter=";"))

        self.assertGreater(len(rows), 5)


class ExportViewTests(TestCase):
    def setUp(self):
        cache.clear()
        self.user = User.objects.create_user(username="eksport", password="haslo12345")
        self.client.force_login(self.user)
        self.audit = Audit.objects.create(
            url="https://example.com", status=Audit.Status.COMPLETED, owner=self.user
        )
        AuditMetric.objects.create(
            audit=self.audit, category="seo", key="title", status="warning", value={"note": "Za krótki."},
        )

    def test_xlsx_download_has_correct_headers(self):
        response = self.client.get(reverse("auditor:export_report", kwargs={"pk": self.audit.pk}), {"format": "xlsx"})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment;", response["Content-Disposition"])
        self.assertIn(".xlsx", response["Content-Disposition"])

    def test_csv_download_is_utf8_with_bom(self):
        response = self.client.get(reverse("auditor:export_report", kwargs={"pk": self.audit.pk}), {"format": "csv"})

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"\xef\xbb\xbf"), "Excel wymaga BOM do rozpoznania UTF-8")

    def test_unknown_format_redirects_with_message(self):
        response = self.client.get(reverse("auditor:export_report", kwargs={"pk": self.audit.pk}), {"format": "pdf"})

        self.assertRedirects(response, reverse("auditor:detail", kwargs={"pk": self.audit.pk}))

    def test_foreign_audit_export_returns_404(self):
        other = User.objects.create_user(username="obcy3", password="haslo12345")
        foreign = Audit.objects.create(url="https://obca.pl", owner=other)

        response = self.client.get(reverse("auditor:export_report", kwargs={"pk": foreign.pk}))

        self.assertEqual(response.status_code, 404)

    def test_sheets_export_requires_connected_google_account(self):
        response = self.client.post(
            reverse("auditor:export_to_google_sheets", kwargs={"pk": self.audit.pk})
        )

        self.assertRedirects(response, reverse("auditor:detail", kwargs={"pk": self.audit.pk}))

    def test_sheets_export_rejects_get(self):
        """Tworzenie arkusza zapisuje dane na koncie użytkownika - musi wymagać POST."""
        response = self.client.get(
            reverse("auditor:export_to_google_sheets", kwargs={"pk": self.audit.pk})
        )

        self.assertRedirects(response, reverse("auditor:detail", kwargs={"pk": self.audit.pk}))

    def test_successful_sheets_export_redirects_to_spreadsheet(self):
        with self.settings(TOKEN_ENCRYPTION_KEY=""):
            self.audit.ga4_refresh_token_encrypted = "token"
            self.audit.save(update_fields=["ga4_refresh_token_encrypted"])

            sheets_service = MagicMock()
            sheets_service.create_report.return_value = "https://docs.google.com/spreadsheets/d/abc"

            with patch("auditor.views._build_credentials_from_refresh_token", return_value=MagicMock()), \
                    patch("auditor.views.GoogleSheetsService", return_value=sheets_service):
                response = self.client.post(
                    reverse("auditor:export_to_google_sheets", kwargs={"pk": self.audit.pk})
                )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response["Location"], "https://docs.google.com/spreadsheets/d/abc")

    def test_missing_scope_shows_reconnect_message(self):
        with self.settings(TOKEN_ENCRYPTION_KEY=""):
            self.audit.ga4_refresh_token_encrypted = "token"
            self.audit.save(update_fields=["ga4_refresh_token_encrypted"])

            sheets_service = MagicMock()
            sheets_service.create_report.side_effect = MissingSheetsScopeError("Połącz konto ponownie.")

            with patch("auditor.views._build_credentials_from_refresh_token", return_value=MagicMock()), \
                    patch("auditor.views.GoogleSheetsService", return_value=sheets_service):
                response = self.client.post(
                    reverse("auditor:export_to_google_sheets", kwargs={"pk": self.audit.pk}), follow=True
                )

        self.assertContains(response, "Połącz konto ponownie.")


class GoogleSheetsServiceTests(SimpleTestCase):
    def test_missing_scope_error_is_translated(self):
        from googleapiclient.errors import HttpError

        response = MagicMock(status=403)
        error = HttpError(resp=response, content=b'{"error": {"message": "insufficient scope"}}')

        translated = GoogleSheetsService()._translate_error(error)

        self.assertIsInstance(translated, MissingSheetsScopeError)

    def test_generic_error_is_translated_to_export_error(self):
        from googleapiclient.errors import HttpError

        response = MagicMock(status=500)
        error = HttpError(resp=response, content=b'{"error": {"message": "server error"}}')

        translated = GoogleSheetsService()._translate_error(error)

        self.assertIsInstance(translated, SheetsExportError)
        self.assertNotIsInstance(translated, MissingSheetsScopeError)


class SitemapEndpointTests(TestCase):
    def setUp(self):
        cache.clear()
        self.user = User.objects.create_user(username="mapa", password="haslo12345")
        self.client.force_login(self.user)

    def test_requires_url_parameter(self):
        response = self.client.get(reverse("auditor:sitemap_suggestions"))

        self.assertEqual(response.status_code, 400)

    def test_returns_suggestions_from_service(self):
        service = MagicMock()
        service.suggest_pages.return_value = {
            "available": True, "sitemap_url": "https://example.com/sitemap.xml",
            "suggestions": {"product": "https://example.com/produkt/x"},
            "scanned_urls": 12, "error": None,
        }

        with patch("auditor.views.SitemapService", return_value=service):
            response = self.client.get(reverse("auditor:sitemap_suggestions"), {"url": "https://example.com"})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["suggestions"]["product"], "https://example.com/produkt/x")

    def test_requires_login(self):
        self.client.logout()

        response = self.client.get(reverse("auditor:sitemap_suggestions"), {"url": "https://example.com"})

        self.assertEqual(response.status_code, 302)
        self.assertIn("/login/", response["Location"])
